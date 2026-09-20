#!/usr/bin/env python3
"""Genera el Excel de revision del censo de virales de medios."""
import csv, json, os
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.expanduser(os.environ.get("MEDIOS_DIR", "~/typesafe-lab/politica/medios"))
rows = list(csv.DictReader(open(os.path.join(BASE, "virales_clasificados.csv"))))
idx = json.load(open(os.path.join(BASE, "sesgo_medios.json")))["por_medio"]

wb = Workbook()
ws = wb.active
ws.title = "Clasificacion"
head = ["Fecha", "Medio", "RTs", "Likes", "Vistas", "Gate", "Relevante", "Partido", "Conf partido",
        "Direccion", "Conf direccion", "Ironia", "Texto", "Enlace"]
ws.append(head)
for c in range(1, len(head) + 1):
    ws.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
    ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor="333333")
PAT = {"PSOE": "DDEBF7", "PP": "FCE4D6", "Vox": "FFF2CC", "Sumar": "E2EFDA"}


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


for r in rows:
    probe = (r["texto"] or "")[:280]
    ws.append([r["fecha"][:10], r["handle"], num(r["retweets"]), num(r["likes"]), num(r["views"]),
               num(r["gate_p"]), r["relevante"], r["partido"], num(r["partido_conf"]),
               r["direccion"], num(r["direccion_conf"]), num(r["ironia_p"]), probe, r["url"]])

widths = [11, 17, 7, 7, 8, 7, 10, 9, 11, 11, 12, 8, 90, 44]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:N{ws.max_row}"
low = PatternFill("solid", fgColor="FFF2CC")
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if row[6].value == "dudoso" or (row[8].value is not None and row[8].value < 0.6) or (row[10].value is not None and row[10].value < 0.6):
        for cell in row:
            cell.fill = low
    part = row[7].value
    if part in PAT:
        row[7].fill = PatternFill("solid", fgColor=PAT[part])

# --- hoja por medio ---
ws2 = wb.create_sheet("Por medio")
ws2.append(["Medio", "Virales", "Con partido", "A la izquierda", "A la derecha", "Neutro", "Indice (der menos izq)"])
for c in range(1, 8):
    ws2.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
    ws2.cell(row=1, column=c).fill = PatternFill("solid", fgColor="333333")
virales = Counter(r["handle"] for r in rows)
tabla = []
for h, v in idx.items():
    claro = v["izq"] + v["der"]
    s = (v["der"] - v["izq"]) / claro if claro else 0
    tabla.append((s, h, v))
for s, h, v in sorted(tabla):
    ws2.append([h, virales.get(h, 0), v["n"], v["izq"], v["der"], v["neutro"], round(s, 3)])
for i, w in enumerate([22, 9, 12, 14, 13, 9, 22], start=1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.freeze_panes = "A2"

# --- hoja top ---
ws3 = wb.create_sheet("Top virales")
ws3.append(["RTs", "Medio", "Partido", "Direccion", "Fecha", "Texto", "Enlace"])
for c in range(1, 8):
    ws3.cell(row=1, column=c).font = Font(bold=True, color="FFFFFF")
    ws3.cell(row=1, column=c).fill = PatternFill("solid", fgColor="333333")
rel = [r for r in rows if r["partido"]]
for r in sorted(rel, key=lambda x: -int(x["retweets"]))[:200]:
    ws3.append([int(r["retweets"]), r["handle"], r["partido"], r["direccion"], r["fecha"][:10],
                (r["texto"] or "")[:250], r["url"]])
for i, w in enumerate([7, 17, 9, 11, 11, 80, 44], start=1):
    ws3.column_dimensions[get_column_letter(i)].width = w

# --- hoja resumen ---
ws4 = wb.create_sheet("Resumen")
res = json.load(open(os.path.join(BASE, "virales_resumen.json")))
izq = sum(v["izq"] for v in idx.values())
der = sum(v["der"] for v in idx.values())
L = [
    ["Censo de tuits virales (>100 RT) de la lista X «Spanish Generalist Media»"],
    ["Ventana", "19 sep 2025 a 19 sep 2026"],
    ["Fuente", "Apify, actor apidojo/twitter-scraper-lite, consultas from:<medio> min_retweets:100 por mes"],
    ["Coste del censo", "21,41 $ (20.404 tuits)"],
    ["Coste de la clasificacion", "10,02 $ (ficha de contexto de Gemini 3.7 Flash con busqueda de Google, 12.520 tuits)"],
    [],
    ["Tuits virales descargados", len(rows)],
    ["Medios con algun viral", len({r["handle"] for r in rows})],
    ["Clasificador", "TypeSafe Jev, dos pasos (gate + partido, direccion, ironia). El paso 2 recibe una ficha de contexto con las personas, empresas y casos del tuit, resuelta con Gemini 3.7 Flash y su busqueda de Google."],
    [],
    ["Paso 1: puede afectar a un partido", ""],
    ["Si", res["gate_si"]], ["Dudoso", res["gate_dudoso"]], ["No", res["gate_no"]],
    [],
    ["Paso 2: partido afectado (sobre los 12.520 que pasan)", ""],
] + [[k, v] for k, v in sorted(res["por_partido"].items(), key=lambda x: -x[1])] + [
    [],
    ["Direccion", ""],
] + [[k, v] for k, v in sorted(res["por_direccion"].items(), key=lambda x: -x[1])] + [
    [],
    ["Cruce partido / direccion", ""],
] + [[k, v] for k, v in sorted(res["cruce"].items(), key=lambda x: -x[1])] + [
    [],
    ["Indice global derecha menos izquierda", round((der - izq) / (der + izq), 3) if (der + izq) else None],
    ["Ironia media detectada", res["ironia_media"]],
    ["Errores de API", res["errores"]],
    [],
    ["Avisos", "Solo se analiza el texto del tuit: los enlaces y las imagenes no se abren."],
    ["", "Los tuits virales no son una muestra neutral de la linea editorial: miden que contenido se comparte mas."],
    ["", "El gate y la direccion son juicios del modelo; la hoja Clasificacion deja ver la confianza de cada fila."],
]
for r in L:
    ws4.append(r)
ws4.column_dimensions["A"].width = 44
ws4.column_dimensions["B"].width = 60
for r in (1, 10, 14, 26):
    try:
        ws4.cell(row=r, column=1).font = Font(bold=True)
    except Exception:
        pass

path = os.path.join(BASE, "medios_virales_clasificacion.xlsx")
wb.save(path)
print("guardado:", path, os.path.getsize(path), "bytes |", ws.max_row - 1, "filas clasificadas")
